"""Retired maintenance commands cannot write or call recognition in managed mode."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import pandas as pd
import build_index
from scripts import apply_main_bank_update as main_update
from scripts import backfill_letter_bank_dimensions as dimensions
from scripts import normalize_excel_load_units as units
from scripts import store_unindexed_questions as unindexed
from scripts import write_symbolic_structure_types as structures
from scripts.build_symbolic_bank import mapped_symbolic_loads


class LegacyWriterTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="lida-retired-writer-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve() / "fixture"
        self.root.mkdir()
        self.book = self.root / "4力法.xlsx"
        self.backups = self.root / "backups"
        self.rel = "4力法/题目/1.png"
        book = Workbook(); sheet = book.active
        sheet.append(["题目名称", "荷载", "结构类型"])
        sheet.append([self.rel, json.dumps({"loads": [{"type": "集中", "raw": "10kN"}]}), "梁"])
        book.save(self.book); book.close()
        self.environment = patch.dict(os.environ, {"TIKU_BANK_STORE": str(self.root / "published")})
        self.environment.start(); self.addCleanup(self.environment.stop)

    def inventory(self):
        return [(p.relative_to(self.root).as_posix(), hashlib.sha256(p.read_bytes()).hexdigest() if p.is_file() else None)
                for p in sorted(self.root.rglob("*"))]

    def test_cli_commands_and_apply_flags_fail_before_inputs_or_side_effects(self):
        source = Path(__file__).resolve().parents[1]
        commands = {"apply_main_bank_update": "--dry-run", "build_symbolic_bank": "--force",
            "store_unindexed_questions": "--apply", "backfill_letter_bank_dimensions": "--dry-run",
            "normalize_excel_load_units": "--apply", "write_symbolic_structure_types": "--apply", "build_index": None}
        before = self.inventory()
        for name, flag in commands.items():
            for options in (([], [flag], ["--help"]) if flag else ([],)):
                with self.subTest(name=name, options=options):
                    entry = source / ("build_index.py" if name == "build_index" else "scripts/" + name + ".py")
                    result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(entry), *options],
                        cwd=self.root, capture_output=True, text=True, encoding="utf-8", timeout=30)
                    if options == ["--help"]:
                        self.assertEqual(result.returncode, 0, result.stderr)
                    else:
                        self.assertNotEqual(result.returncode, 0)
                        self.assertIn("Published banks are read-only here", result.stderr)
                    self.assertEqual(self.inventory(), before)

    def test_direct_write_functions_fail_without_touching_workbooks_or_dataframes(self):
        frame = pd.DataFrame([{"题目名称": self.rel, "结构类型": "梁"}])
        before_frame = frame.copy(deep=True)
        operations = [lambda: build_index.process_chapter(None, self.root),
            lambda: build_index.save_chapter_excel("4力法", []),
            lambda: main_update.apply_chapter_update(self.book, {self.rel}, [], False),
            lambda: dimensions.backfill_file(self.book, {self.rel: "8m×4m"}, dry_run=False),
            lambda: units.normalize_workbook(self.book, dry_run=False, backup_dir=self.backups),
            lambda: unindexed.append_plans_to_workbook(self.book, [], bank="main", dry_run=False, backup_dir=self.backups),
            lambda: unindexed.apply_ready_plans([], dry_run=False, backup_dir=self.backups),
            lambda: structures.write_workbook(self.book, frame, [{"rel_path": self.rel, "structure_type": "钢架"}], self.backups)]
        before = self.inventory()
        for operation in operations:
            with self.assertRaisesRegex(PermissionError, "read-only"):
                operation()
            self.assertEqual(before, self.inventory())
            pd.testing.assert_frame_equal(before_frame, frame)

    def test_backup_helpers_cannot_create_uncontrolled_outputs(self):
        before = self.inventory()
        for operation in (lambda: dimensions.backup_bank(self.root),
            lambda: units.backup_once(self.book, self.backups, set()),
            lambda: unindexed.backup_workbook(self.book, self.backups, "main"),
            lambda: structures.backup_workbook(self.book, self.backups)):
            with self.assertRaisesRegex(PermissionError, "read-only"):
                operation()
            self.assertEqual(before, self.inventory())

    def test_read_only_helpers_preserve_bytes_and_dry_run_creates_no_directory(self):
        before = self.inventory()
        main_update.apply_chapter_update(self.book, {self.rel}, [], True)
        dimensions.backfill_file(self.book, {self.rel: "8m×4m"}, dry_run=True)
        units.normalize_workbook(self.book, dry_run=True, backup_dir=self.backups)
        plan = SimpleNamespace(rel_path=self.rel)
        result = unindexed.append_plans_to_workbook(self.root / "absent" / "new.xlsx", [plan], bank="main", dry_run=True, backup_dir=self.backups)
        self.assertEqual(result.appended, [self.rel])
        self.assertEqual(before, self.inventory())

    def test_explicit_legacy_environment_still_performs_real_updates_on_temporary_bank(self):
        os.environ.pop("TIKU_BANK_STORE")
        dimensions.backfill_file(self.book, {self.rel: "8m×4m"}, dry_run=False)
        report = units.normalize_workbook(self.book, dry_run=False, backup_dir=self.backups)
        self.assertGreater(report["changed"], 0)
        frame = pd.read_excel(self.book, keep_default_na=False)
        structures.write_workbook(self.book, frame, [{"rel_path": self.rel, "structure_type": "钢架"}], self.backups)
        book = load_workbook(self.book)
        try:
            values = {cell.value: book.active.cell(2, cell.column).value for cell in book.active[1]}
            self.assertEqual(values["结构类型"], "钢架")
            self.assertEqual(values["长×宽"], "8m×4m")
        finally:
            book.close()
        main_update.apply_chapter_update(self.book, {self.rel}, [], False)
        book = load_workbook(self.book)
        try:
            self.assertEqual(book.active.max_row, 1)
        finally:
            book.close()

    def test_pure_conversion_remains_available_to_managed_recognition(self):
        self.assertTrue(mapped_symbolic_loads([{"type": "集中", "raw": "P"}]))
        self.assertEqual(main_update.strip_load_unit("10kN"), "10")


if __name__ == "__main__":
    unittest.main()
